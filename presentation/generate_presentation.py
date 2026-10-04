"""Build the jury deck, screenshot checklist and static layout previews."""

from __future__ import annotations

import io
import json
import math
from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation"
SCREENSHOTS = OUT / "assets" / "screenshots"
PREVIEWS = OUT / "previews"
PPTX = OUT / "BiKing_HackYeah.pptx"
W, H = 16.0, 9.0
M = 0.72
FONT = "Arial"
BG = "F5F7F2"
WHITE = "FFFFFF"
INK = "14231D"
DARK = "10231B"
PANEL = "1B3328"
GREEN = "087B45"
LIME = "65E5A0"
MINT = "E4F3E7"
MUTED = "607068"
PALE = "C0D0C5"
BORDER = "D9E2D9"
GOLD = "F3C76A"
SHADOW = "E9EEE6"
EMU = 914400
SCALE = 100
FONT_DIR = Path("/usr/share/fonts/truetype/liberation")

SCREENSHOT_SPECS = [
    {
        "slide": 3,
        "filename": "main_ui.png",
        "title": "Main application screen",
        "visible": "Kraków map, A/B waypoints, scenario, profile and departure controls.",
        "state": "Use the real Flask/Leaflet app in light mode at a desktop viewport. "
        "Select the heatwave scenario and a profile; keep the map and controls unobstructed. "
        "Choose origin and destination, then let route calculation finish.",
        "exclude": "Loading overlay, browser chrome, devtools, errors, unrelated tabs, "
        "mock data passed off as city data, open accessibility popovers.",
        "why": "Shows that the user journey lives in one working application screen.",
        "size": "1600 × 950 or a similar landscape aspect ratio; capture the application only.",
    },
    {
        "slide": 7,
        "filename": "route_comparison.png",
        "title": "Fastest vs More comfortable",
        "visible": "Both routes on one map, identical origin/destination, dashed grey Fastest "
        "and solid green More comfortable, route legend and visible route separation.",
        "state": "First load matching city graph and shade artifacts. Verify /api/health reports "
        "real graph, shade, environment and exposure modules. Use Heatwave · 3 Jul 2025, "
        "14:00 and Senior / Child as a starting point. The existing Load demo route preset "
        "uses Kazimierz → Rondo Mogilskie → Nowa Huta. Check that the returned routes "
        "actually differ; otherwise choose another real A/B pair. Keep all factors enabled. "
        "Capture the map and save the /api/route response from that same run for metric entry.",
        "exclude": "Different endpoints or times for the two routes, synthetic fixture networks, "
        "a mock result labelled real, manual route drawings, invented gains, loading states. "
        "Retain OpenStreetMap attribution when map tiles are visible.",
        "why": "Demonstrates that environmental edge costs can change the route geometry.",
        "size": "1900 × 850 or a similar wide map crop; leave enough detail to see both routes.",
    },
    {
        "slide": 8,
        "filename": "route_details.png",
        "title": "Route details and trade-off controls",
        "visible": "Real route cards, time/distance, shade, felt temperature, PM2.5 estimate, "
        "poor-air/high-UV minutes, actual avoided-street explanation, comfort/time slider "
        "and GPX export control. Keep some map visible if practical.",
        "state": "Use the same fully calculated route request as slide 7, preferably with "
        "more than one valid trade-off option. Select More comfortable, scroll only enough "
        "to show its card and the trade-off panel. If the UI reports no useful detour, "
        "show that honest state or select a different real demo pair.",
        "exclude": "Mock numbers, fabricated explanation text, clipped cards, open tooltips "
        "obscuring metrics, health claims beyond the model, debug panes and loading overlays.",
        "why": "Shows how GIS and environmental modelling become an understandable choice.",
        "size": "1500 × 850 or a similar landscape crop. Preserve text legibility.",
    },
]


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def shape(
    slide,
    x,
    y,
    w,
    h,
    fill=None,
    stroke=None,
    radius=False,
    kind=None,
    line_width=1.0,
    name="shape",
):
    preset = kind or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    obj = slide.shapes.add_shape(preset, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.name = name
    if radius:
        obj.adjustments[0] = 0.12
    if fill:
        obj.fill.solid()
        obj.fill.fore_color.rgb = rgb(fill)
    else:
        obj.fill.background()
    if stroke:
        obj.line.color.rgb = rgb(stroke)
        obj.line.width = Pt(line_width)
    else:
        obj.line.fill.background()
    return obj


def text(
    slide,
    value,
    x,
    y,
    w,
    h,
    size=22,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    name="text",
):
    obj = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    obj.name = name
    frame = obj.text_frame
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = Pt(size * 1.15)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return obj


def line(
    slide,
    x1,
    y1,
    x2,
    y2,
    color=BORDER,
    width=1.5,
    dashed=False,
    arrow=False,
    name="line",
):
    obj = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    obj.name = name
    obj.line.color.rgb = rgb(color)
    obj.line.width = Pt(width)
    if dashed:
        dash = OxmlElement("a:prstDash")
        dash.set("val", "dash")
        obj.line._get_or_add_ln().append(dash)
    if arrow:
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        end.set("w", "sm")
        end.set("len", "sm")
        obj.line._get_or_add_ln().append(end)
    return obj


def polygon(slide, points, fill, name="polygon"):
    coordinates = [(int(x * EMU), int(y * EMU)) for x, y in points]
    builder = slide.shapes.build_freeform(*coordinates[0])
    builder.add_line_segments(coordinates[1:], close=True)
    obj = builder.convert_to_shape()
    obj.name = name
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(fill)
    obj.line.fill.background()
    return obj


def card(slide, x, y, w, h, dark=False, accent=False):
    if not dark:
        shape(slide, x + 0.04, y + 0.06, w, h, SHADOW, radius=True, name="card-shadow")
    return shape(
        slide,
        x,
        y,
        w,
        h,
        PANEL if dark else (MINT if accent else WHITE),
        "355141" if dark else BORDER,
        radius=True,
        name="card",
    )


def chip(slide, label, x, y, w, dark=False):
    shape(slide, x, y, w, 0.43, PANEL if dark else MINT, radius=True)
    text(
        slide,
        label,
        x + 0.12,
        y + 0.095,
        w - 0.24,
        0.28,
        15.5,
        LIME if dark else GREEN,
        bold=True,
    )


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def footer(slide, number, dark=False):
    col = PALE if dark else MUTED
    line(slide, M, 8.48, W - M, 8.48, "2B4435" if dark else BORDER, 0.8)
    text(
        slide,
        "BiKing  /  HackYeah 2026",
        M,
        8.64,
        6,
        0.21,
        12.5,
        col,
        name="footer-brand",
    )
    text(
        slide,
        f"{number:02d} / 10",
        14.05,
        8.62,
        1.23,
        0.24,
        13,
        col,
        align=PP_ALIGN.RIGHT,
        name="footer-number",
    )


def new_slide(
    prs, number, kicker, title=None, subtitle=None, dark=False, title_size=43
):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(DARK if dark else BG)
    shape(slide, M, 0.46, 0.38, 0.07, LIME if dark else GREEN)
    text(
        slide,
        kicker.upper(),
        M + 0.55,
        0.39,
        13.8,
        0.32,
        15.5,
        LIME if dark else GREEN,
        bold=True,
    )
    if title:
        lines = title.count("\n") + 1
        height = lines * title_size * 1.15 / 72 + 0.07
        text(
            slide,
            title,
            M,
            0.95,
            14.56,
            height,
            title_size,
            WHITE if dark else INK,
            bold=True,
            name="slide-title",
        )
    if subtitle:
        text(slide, subtitle, M, 1.82, 14.5, 0.46, 21, PALE if dark else MUTED)
    footer(slide, number, dark)
    return slide


def picture_contained(slide, path, x, y, w, h, name="picture"):
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    obj = slide.shapes.add_picture(
        str(path),
        Inches(x + (w - pw) / 2),
        Inches(y + (h - ph) / 2),
        width=Inches(pw),
        height=Inches(ph),
    )
    obj.name = name
    obj._element.nvPicPr.cNvPr.set("descr", path.name)
    return obj


def screenshot(slide, filename, title, description, x, y, w, h):
    card(slide, x, y, w, h)
    path = SCREENSHOTS / filename
    if path.is_file():
        picture_contained(
            slide,
            path,
            x + 0.12,
            y + 0.12,
            w - 0.24,
            h - 0.24,
            name=f"screenshot-{filename}",
        )
        return
    obj = shape(
        slide,
        x + 0.18,
        y + 0.18,
        w - 0.36,
        h - 0.36,
        BG,
        BORDER,
        radius=True,
        name=f"placeholder-{filename}",
    )
    dash = OxmlElement("a:prstDash")
    dash.set("val", "dash")
    obj.line._get_or_add_ln().append(dash)
    cy = y + h / 2
    shape(
        slide,
        x + w / 2 - 0.23,
        cy - 1.16,
        0.46,
        0.34,
        None,
        GREEN,
        radius=True,
        line_width=1.6,
    )
    shape(
        slide,
        x + w / 2 - 0.075,
        cy - 1.065,
        0.15,
        0.15,
        None,
        GREEN,
        kind=MSO_SHAPE.OVAL,
        line_width=1.3,
    )
    text(
        slide,
        "SCREENSHOT NEEDED",
        x + 0.35,
        cy - 0.56,
        w - 0.7,
        0.42,
        23,
        GREEN,
        True,
        PP_ALIGN.CENTER,
    )
    text(
        slide, title, x + 0.35, cy + 0.04, w - 0.7, 0.43, 22, INK, True, PP_ALIGN.CENTER
    )
    text(
        slide,
        description,
        x + 0.35,
        cy + 0.59,
        w - 0.7,
        0.84,
        17,
        MUTED,
        align=PP_ALIGN.CENTER,
    )
    text(
        slide,
        f"Expected: {filename}",
        x + 0.35,
        y + h - 0.67,
        w - 0.7,
        0.35,
        16,
        MUTED,
        align=PP_ALIGN.CENTER,
    )


def sun_icon(slide, x, y, size=0.7, color=GOLD):
    cx, cy = x + size / 2, y + size / 2
    shape(
        slide,
        x + size * 0.23,
        y + size * 0.23,
        size * 0.54,
        size * 0.54,
        color,
        kind=MSO_SHAPE.OVAL,
    )
    for i in range(8):
        angle = i * math.pi / 4
        line(
            slide,
            cx + math.cos(angle) * size * 0.37,
            cy + math.sin(angle) * size * 0.37,
            cx + math.cos(angle) * size * 0.49,
            cy + math.sin(angle) * size * 0.49,
            color,
            1.8,
        )


def bicycle(slide, x, y, scale=1.0, color=GREEN):
    for wheel_x in (x, x + 0.72 * scale):
        shape(
            slide,
            wheel_x,
            y + 0.33 * scale,
            0.4 * scale,
            0.4 * scale,
            None,
            color,
            kind=MSO_SHAPE.OVAL,
            line_width=1.7,
        )
    for a, b in [
        ((0.2, 0.53), (0.48, 0.22)),
        ((0.48, 0.22), (0.72, 0.53)),
        ((0.72, 0.53), (0.2, 0.53)),
        ((0.48, 0.22), (0.84, 0.22)),
        ((0.84, 0.22), (0.92, 0.53)),
        ((0.84, 0.22), (0.8, 0.09)),
    ]:
        line(
            slide,
            x + a[0] * scale,
            y + a[1] * scale,
            x + b[0] * scale,
            y + b[1] * scale,
            color,
            1.7,
        )
    shape(
        slide,
        x + 0.57 * scale,
        y - 0.19 * scale,
        0.16 * scale,
        0.16 * scale,
        color,
        kind=MSO_SHAPE.OVAL,
    )
    for a, b in [
        ((0.63, 0), (0.47, 0.16)),
        ((0.63, 0), (0.82, 0.14)),
        ((0.47, 0.16), (0.62, 0.34)),
        ((0.62, 0.34), (0.56, 0.48)),
    ]:
        line(
            slide,
            x + a[0] * scale,
            y + a[1] * scale,
            x + b[0] * scale,
            y + b[1] * scale,
            color,
            2.0,
        )


def make_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = (
        "BiKing — The shortest route is not always the best route"
    )
    prs.core_properties.subject = (
        "HackYeah jury presentation: Sport & Healthcare and Smart City"
    )
    prs.core_properties.author = "BiKing"
    prs.core_properties.keywords = (
        "cycling, shade, environmental routing, Kraków, HackYeah"
    )

    s = new_slide(prs, 1, "HackYeah 2026 / Kraków", dark=True)
    text(s, "BiKing", M, 1.12, 9, 1.1, 68, WHITE, True)
    text(s, "The shortest route", M, 2.75, 9.3, 0.8, 43, WHITE, True)
    text(s, "is not always", M, 3.48, 9.3, 0.8, 43, WHITE, True)
    text(s, "the best route.", M, 4.21, 9.3, 0.8, 43, LIME, True)
    text(
        s,
        "Cycling routes adapted to\nthe city around you.",
        M,
        5.57,
        9.2,
        1.03,
        27,
        PALE,
    )
    chip(s, "SPORT & HEALTHCARE", M, 7.26, 3.33, True)
    chip(s, "SMART CITY", 4.25, 7.26, 2.0, True)
    shape(s, 10.56, 1.08, 4.72, 6.88, WHITE, radius=True)
    picture_contained(
        s,
        ROOT / "app/static/logo.png",
        10.91,
        1.61,
        4.0,
        4.0,
        name="BiKing-original-logo",
    )
    text(s, "CITY CONTEXT", 10.87, 6.0, 4.08, 0.32, 17, GREEN, True, PP_ALIGN.CENTER)
    text(
        s,
        "Air  ·  Heat  ·  UV  ·  Shade",
        10.87,
        6.55,
        4.08,
        0.43,
        20,
        INK,
        align=PP_ALIGN.CENTER,
    )
    line(s, 11.72, 7.35, 14.12, 7.35, GREEN, 3.0, arrow=True)
    notes(
        s,
        "Opening: BiKing makes city and environmental context part of bicycle route selection. "
        "The current product is the Flask/Leaflet app, not the obsolete Streamlit demo. "
        "Original logo: app/static/logo.png. No team roster was found, so no names were invented. "
        "This deck serves both requested categories. No medical outcome is claimed.",
    )

    s = new_slide(
        prs,
        2,
        "The problem",
        "Two routes. Two different rides.",
        "Similar distance and time can hide very different conditions.",
    )
    card(s, M, 2.56, 6.75, 4.47)
    card(s, 7.8, 2.56, 7.48, 4.47, accent=True)
    text(s, "THE USUAL COMPARISON", 1.06, 2.9, 6.05, 0.38, 18, MUTED, True)
    text(s, "How far?\nHow long?", 1.07, 3.67, 5.95, 1.67, 43, INK, True)
    text(s, "Distance and travel time", 1.07, 6.06, 5.95, 0.45, 23, MUTED)
    text(s, "THE CONDITIONS OF THE RIDE", 8.15, 2.9, 6.7, 0.38, 18, GREEN, True)
    factors = [
        ("Air", "PM2.5 · PM10 · NO₂"),
        ("Heat", "Felt temperature / UTCI"),
        ("UV", "Shade-adjusted UV"),
        ("Shade", "Buildings + trees"),
    ]
    for i, (title, desc) in enumerate(factors):
        x = 8.15 + (i % 2) * 3.42
        y = 3.64 + (i // 2) * 1.4
        shape(s, x, y, 3.08, 1.16, WHITE, radius=True)
        text(s, title, x + 0.2, y + 0.15, 2.69, 0.48, 27, GREEN, True)
        text(s, desc, x + 0.2, y + 0.74, 2.7, 0.32, 16.5, MUTED)
    text(
        s,
        "A route is also an environmental exposure.",
        M,
        7.49,
        14.56,
        0.6,
        29,
        INK,
        True,
    )
    notes(
        s,
        "This is a conceptual comparison, not a claim that a particular measured pair has "
        "similar time or a given exposure reduction. Implemented factors: PM2.5, PM10, NO₂, "
        "UTCI/felt temperature, shade-adjusted UV, buildings and trees. Sources: "
        "app/env/exposure.py:compute_edge_exposure; app/engine/metrics.py:route_metrics. "
        "Pollution and thermal fields are model estimates, not sensors on every street. "
        "Shade does not guarantee cleaner air; the factors remain distinct.",
    )

    s = new_slide(prs, 3, "The product", "From A to B, with context.")
    steps = [
        ("Choose points", "Start, destination\n+ optional via points"),
        ("Set context", "Scenario, departure time\nand rider profile"),
        ("Compare routes", "Fastest vs\nMore comfortable"),
        ("Adjust the trade-off", "Choose how much extra\ntime comfort is worth"),
        ("Ride", "Export GPX to\nyour cycling app"),
    ]
    for i, (title, desc) in enumerate(steps):
        y = 2.25 + i * 1.13
        if i < 4:
            line(s, 0.955, y + 0.43, 0.955, y + 1.1, BORDER, 2)
        shape(s, M, y, 0.47, 0.47, GREEN, kind=MSO_SHAPE.OVAL)
        text(s, str(i + 1), M, y + 0.075, 0.47, 0.31, 18, WHITE, True, PP_ALIGN.CENTER)
        text(s, title, 1.43, y + 0.01, 3.7, 0.42, 23, INK, True)
        text(s, desc, 1.43, y + 0.48, 3.7, 0.7, 18, MUTED)
    screenshot(
        s,
        "main_ui.png",
        "Final main application screen",
        "Map visible · origin / destination\nRoute controls · clean application state",
        5.35,
        2.17,
        9.93,
        5.81,
    )
    notes(
        s,
        "UI source: app/static/index.html and app/static/app.js. Users click or drag map "
        "waypoints; the route automatically recalculates. Up to five points are accepted. "
        "Scenario/profile/departure controls and selected heat/air/UV factors are implemented. "
        "POST /api/route/tradeoff powers the comfort/time slider. GPX export is implemented "
        "in routeToGpx/exportGpx. Export is not built-in turn-by-turn navigation. Screenshot "
        "placeholder is deliberate and is replaced automatically if main_ui.png exists.",
    )

    s = new_slide(
        prs,
        4,
        "The system",
        "The city becomes part of the routing model.",
        title_size=38,
    )
    inputs = [
        (M, "ENVIRONMENT", "Open-Meteo + GIOŚ", "Weather · UV\nPM2.5 · PM10 · NO₂"),
        (
            5.68,
            "CITY GEOMETRY",
            "GUGiK LoD1 / OSM / MSIP",
            "Footprints · building heights\nTree / green-cover geometry",
        ),
        (
            10.64,
            "BICYCLE NETWORK",
            "OpenStreetMap",
            "pyosmium → OSMnx\nDirected streets · road classes",
        ),
    ]
    for x, label, source, detail in inputs:
        card(s, x, 2.31, 4.64, 1.91)
        text(s, label, x + 0.24, 2.57, 4.17, 0.37, 18, GREEN, True)
        text(s, source, x + 0.24, 3.04, 4.17, 0.4, 19.5, INK, True)
        text(s, detail, x + 0.24, 3.58, 4.17, 0.69, 18, MUTED)
        cx = x + 2.32
        line(s, cx, 4.22, cx, 4.55, GREEN, 1.7)
    line(s, 3.04, 4.55, 12.96, 4.55, GREEN, 1.7)
    line(s, 5.34, 4.55, 5.34, 4.85, GREEN, 1.7, arrow=True)
    card(s, M, 4.89, 9.21, 1.42, accent=True)
    text(s, "PER-EDGE EXPOSURE MODEL", 1.02, 5.11, 8.61, 0.42, 23, GREEN, True)
    text(
        s,
        "Shade → UTCI + air index + effective UV → fuzzy discomfort",
        1.02,
        5.7,
        8.61,
        0.42,
        19,
        INK,
    )
    card(s, 10.64, 4.89, 4.64, 1.42)
    text(s, "USER CONTEXT", 10.91, 5.13, 4.1, 0.4, 20, GREEN, True)
    text(s, "Time · profile · selected factors", 10.91, 5.73, 4.1, 0.38, 18.5, MUTED)
    line(s, 12.96, 6.31, 12.96, 6.61, GREEN, 1.7)
    line(s, 12.96, 6.61, 2.84, 6.61, GREEN, 1.7)
    line(s, 2.84, 6.31, 2.84, 6.93, GREEN, 1.7, arrow=True)
    for x, label in [
        (M, "Road-segment costs"),
        (5.91, "Dijkstra route search"),
        (11.1, "BiKing route choices"),
    ]:
        shape(
            s,
            x,
            6.96,
            4.18,
            0.96,
            GREEN if x == 11.1 else WHITE,
            None if x == 11.1 else BORDER,
            radius=True,
        )
        text(
            s,
            label,
            x + 0.14,
            7.23,
            3.9,
            0.44,
            21,
            WHITE if x == 11.1 else INK,
            True,
            PP_ALIGN.CENTER,
        )
    line(s, 4.96, 7.44, 5.81, 7.44, GREEN, 2, arrow=True)
    line(s, 10.15, 7.44, 11.0, 7.44, GREEN, 2, arrow=True)
    notes(
        s,
        "Implemented data paths: app/env/fetch.py queries Open-Meteo weather, air-quality "
        "(CAMS) and GIOŚ. GIOŚ supplies a city correction and an IDW spatial background; "
        "road-class factors are calibrated in pipeline/calibrate_road.py. Geometry adapters "
        "in pipeline/p03_buildings.py and pipeline/shade_build.py handle GUGiK LoD1 2024, "
        "OSM and MSIP 2015 green cover, with documented fallbacks. Building heights may be "
        "measured, tagged, inferred from floors, or defaults; tree heights are assumptions. "
        "Graph construction is offline OSM PBF → pyosmium → OSMnx → SciPy CSR. "
        "No Airly, Sentinel or Google Maps integration is claimed. Optional Landsat "
        "heat-map code exists but its artifact is absent here, so it is not a core data-flow claim.",
    )

    s = new_slide(
        prs, 5, "The innovation / edge weights", "Routing beyond distance.", dark=True
    )
    card(s, M, 2.24, 5.04, 2.21, dark=True)
    card(s, 6.11, 2.24, 9.17, 2.21, dark=True)
    text(s, "FASTEST", 1.03, 2.57, 4.41, 0.39, 19, PALE, True)
    text(s, "Edge cost = travel time", 1.03, 3.22, 4.44, 0.52, 26, WHITE, True)
    text(s, "Minimise total ride time.", 1.03, 3.96, 4.44, 0.4, 19, PALE)
    text(s, "BIKING / MORE COMFORTABLE", 6.48, 2.57, 8.43, 0.39, 19, LIME, True)
    text(s, "C = t × (1 + αX + wP)", 6.48, 3.15, 8.43, 0.65, 35, WHITE, True)
    text(
        s,
        "Time × a profile-specific environmental penalty",
        6.48,
        3.99,
        8.43,
        0.38,
        19,
        PALE,
    )
    text(s, "PENALISE THE AVOIDABLE PART", M, 4.97, 9.2, 0.38, 18, LIME, True)
    text(
        s, "X = clip((D − D₁₀) / (1 − D₁₀), 0, 1)", M, 5.59, 9.35, 0.59, 27, WHITE, True
    )
    text(
        s,
        "D: fuzzy discomfort    D₁₀: today's city baseline",
        M,
        6.3,
        9.22,
        0.42,
        20,
        PALE,
    )
    text(
        s,
        "α / w: rider weights    P: relative air + hotspot penalty",
        M,
        6.84,
        9.22,
        0.37,
        18,
        PALE,
    )
    line(s, 10.05, 4.98, 10.05, 6.96, "355141", 1)
    text(
        s, "Same graph.\nDifferent edge costs.", 10.55, 5.15, 4.71, 1.06, 29, LIME, True
    )
    text(
        s,
        "SciPy Dijkstra\non a directed CSR graph",
        10.55,
        6.46,
        4.71,
        0.69,
        19.5,
        PALE,
    )
    text(
        s,
        "Weights are applied before path search.",
        M,
        7.42,
        14.56,
        0.58,
        28,
        WHITE,
        True,
    )
    notes(
        s,
        "Actual formula: app/engine/variants.py:eco_cost. C_e = t_e * "
        "(1 + alpha * X_e + w * P_e). D10 is the 10th percentile of current discomfort "
        "over every graph edge. X = clip((D-D10)/max(1-D10, 1e-6), 0, 1). "
        "P = A/median(A) + clip(A-A10, 0, 1), with A cleaned/clipped to [0.1,6]. "
        "w = 6 for Asthma and 0 for other profiles; alpha is 5/5/4/3 for "
        "standard/asthma/senior/athlete. Profile speeds set t = length/speed. "
        "D comes from profile-specific Mamdani fuzzy LUTs with 11 rules. "
        "The time/comfort slider scales both environmental weights. "
        "app/graph/routing.py uses scipy.sparse.csgraph.dijkstra, not A* or NetworkX "
        "at request time. Summed edge costs select the path; this is not post-hoc colouring. "
        "The displayed baseline expression omits only the 1e-6 division guard for clarity.",
    )

    s = new_slide(
        prs,
        6,
        "The technical highlight / physical shade",
        "We actually calculate the shadows.",
        title_size=42,
    )
    card(s, M, 2.22, 8.9, 4.72)
    text(s, "SCHEMATIC / ROAD CROSS-SECTION", 1.02, 2.5, 8.3, 0.35, 16.5, MUTED, True)
    sun_icon(s, 8.2, 3.12, 0.84)
    text(s, "Sun position", 7.62, 4.1, 1.63, 0.35, 16, GREEN)
    shape(s, 1.17, 6.12, 7.94, 0.34, "DCE5DB", radius=True)
    polygon(
        s,
        [(1.42, 6.13), (5.33, 6.13), (5.78, 6.44), (1.42, 6.44)],
        "B4DFC0",
        "projected-shadow-schematic",
    )
    line(s, 2.02, 5.77, 8.61, 3.54, GOLD, 2.4, dashed=True, arrow=True)
    text(s, "Trace toward the sun", 1.11, 3.27, 4.2, 0.43, 21, GREEN, True)
    text(s, "Ray blocked by the building", 1.11, 3.85, 4.25, 0.41, 19, MUTED)
    shape(s, 5.28, 4.35, 1.35, 1.78, "47614F", radius=True)
    shape(s, 5.49, 4.59, 0.24, 0.31, "BED3C3")
    shape(s, 6.0, 4.59, 0.24, 0.31, "BED3C3")
    shape(s, 5.49, 5.17, 0.24, 0.31, "BED3C3")
    shape(s, 6.0, 5.17, 0.24, 0.31, "BED3C3")
    text(s, "Building / H", 5.11, 6.59, 2.0, 0.34, 17, MUTED)
    bicycle(s, 1.58, 5.54, 0.72)
    text(s, "1.5 m eye height", 1.08, 5.0, 2.87, 0.37, 18, INK)
    for x in (1.63, 2.59, 3.55, 4.51, 7.39, 8.35):
        shape(s, x, 6.22, 0.13, 0.13, GREEN if x < 5.28 else GOLD, kind=MSO_SHAPE.OVAL)
    text(s, "Shaded road samples", 2.43, 6.58, 2.66, 0.35, 17, GREEN)
    card(s, 10.01, 2.22, 5.27, 1.69, accent=True)
    text(s, "A POINT IS SHADED WHEN", 10.28, 2.49, 4.73, 0.36, 17, GREEN, True)
    text(s, "H(d) > 1.5 m + d · tan(θ)", 10.28, 3.05, 4.73, 0.5, 23, INK, True)
    text(s, "θ = solar elevation; d = ray distance", 10.28, 3.62, 4.73, 0.32, 16, MUTED)
    flow = [
        ("01", "Solar geometry", "pvlib elevation + azimuth"),
        ("02", "City geometry", "Building + tree height raster"),
        ("03", "Shade fraction", "Shaded / total road samples"),
        ("04", "Routing effect", "UTCI + UV → fuzzy cost → path"),
    ]
    for i, (num, title, detail) in enumerate(flow):
        y = 4.17 + i * 0.7
        text(s, num, 10.04, y + 0.02, 0.48, 0.36, 18, GREEN, True)
        text(s, title, 10.77, y, 4.5, 0.38, 21, INK, True)
        text(s, detail, 10.77, y + 0.38, 4.5, 0.33, 17, MUTED)
    text(
        s,
        "Precompute 16 × 8 solar bins → look up shade at departure time",
        M,
        7.32,
        14.56,
        0.54,
        25,
        GREEN,
        True,
    )
    text(
        s,
        "2 m raster / ray steps  ·  120 m ray range  ·  road samples at most 15 m apart",
        M,
        7.97,
        14.56,
        0.33,
        17,
        MUTED,
    )
    notes(
        s,
        "Hero implementation: pipeline/p04_height_raster.py rasterizes footprints and "
        "heights in EPSG:2180 at 2 m. pipeline/p05_shade.py:shaded_points traces road "
        "sample rays toward the sun: obstacle height > EYE + d*tan(elevation), "
        "EYE=1.5 m, d=2..120 m in 2 m steps. Points inside obstacles of at least 3 m "
        "are also marked shaded. Equal-length edge cells have midpoint samples no farther "
        "than 15 m apart, minimum two. Their shaded fraction is stored as uint8 in "
        "(E,16,8), indexed by eid. These are model settings, not measured performance. "
        "app/shade/sun.py uses pvlib NREL apparent elevation and azimuth. "
        "app/shade/model.py uses nearest circular azimuth bin, linear elevation "
        "interpolation; daylight elevation is clamped to 5..65 degrees, night returns 1. "
        "Runtime uses prepared tables, not fresh GIS ray-marching per request. "
        "The native diagram is explanatory geometry, not a fabricated app screenshot or "
        "Kraków measurement. No polygon-shadow generation, full 3D ray tracing, "
        "terrain effects or measured tree heights are claimed.",
    )

    s = new_slide(
        prs, 7, "The proof / final demo capture", "Same city. Different route."
    )
    line(s, M + 0.03, 2.22, M + 0.68, 2.22, MUTED, 3, dashed=True)
    text(s, "Fastest", 1.58, 2.02, 2.03, 0.44, 20, MUTED, True)
    line(s, 3.63, 2.22, 4.29, 2.22, GREEN, 3)
    text(s, "More comfortable", 4.48, 2.02, 5.85, 0.44, 20, GREEN, True)
    screenshot(
        s,
        "route_comparison.png",
        "Final route comparison",
        "Same origin + destination\nFastest + environmental route visible",
        M,
        2.63,
        10.36,
        4.64,
    )
    metric_specs = [
        ("Travel time", "min"),
        ("Shade", "% of ride"),
        ("Felt temperature", "°C / UTCI model"),
        ("PM2.5 dose", "µg / model estimate"),
    ]
    for i, (label, unit) in enumerate(metric_specs):
        y = 2.06 + i * 1.32
        card(s, 11.42, y, 3.86, 1.18, accent=i == 1)
        text(s, label, 11.65, y + 0.12, 3.38, 0.35, 19, INK, True)
        text(
            s, "Fastest  —    BiKing  —", 11.65, y + 0.48, 3.38, 0.38, 18.5, GREEN, True
        )
        text(s, unit, 11.65, y + 0.9, 3.38, 0.29, 15.5, MUTED)
    text(
        s,
        "Same departure time, scenario and rider profile.",
        M,
        7.48,
        10.36,
        0.37,
        18,
        MUTED,
    )
    text(s, "Final demo values pending", 11.43, 7.47, 3.84, 0.35, 16, MUTED)
    text(
        s,
        "Environmental conditions can change the path itself.",
        M,
        7.99,
        14.56,
        0.48,
        25,
        GREEN,
        True,
    )
    notes(
        s,
        "This screenshot and the metric values are pending real-city capture. "
        "BiKing here means More comfortable. The metrics are real implemented fields in "
        "app/engine/metrics.py: time_min, shade_pct, utci_avg_c, pm25_dose_ug. "
        "Shade and UTCI are time-weighted; PM2.5 dose is concentration × profile ventilation "
        "× travel time. This is a model estimate, not a personal dosimeter reading. "
        "The checkout lacks graph.npz, edge_coords.npz, edges.parquet and city shade artifacts "
        "under data/processed. Do not use the synthetic shade_preview map or test networks "
        "as Kraków proof. Tests/test_integration.py verifies differing geometries with "
        "the real implementations on fixture graphs, not city performance. "
        "The Fastest route is a BiKing baseline at profile speed, not a Google Maps benchmark. "
        "If both routes coincide, the app explicitly reports that rather than forcing a detour. "
        "Replacing the screenshot never auto-imports numbers; enter matched /api/route values "
        "in the editable PowerPoint cards after the final capture. No reduction is invented.",
    )

    s = new_slide(
        prs,
        8,
        "The experience",
        "Complexity under the hood.\nSimple for the cyclist.",
        title_size=40,
    )
    text(
        s, "Which route fits my\nconditions today?", M, 2.7, 5.13, 1.24, 29, GREEN, True
    )
    benefits = [
        ("Understand the trade-off", "Time, shade, air, heat and UV"),
        ("See the reason", "Avoided streets + segment colours"),
        ("Make it yours", "Comfort slider + GPX export"),
    ]
    for i, (title, detail) in enumerate(benefits):
        y = 4.46 + i * 0.99
        shape(s, M, y + 0.12, 0.1, 0.1, GREEN, kind=MSO_SHAPE.OVAL)
        text(s, title, 1.0, y, 4.76, 0.45, 23, INK, True)
        text(s, detail, 1.0, y + 0.52, 4.77, 0.35, 18, MUTED)
    chip(s, "ACCESSIBLE COLOURS + TEXT", M, 7.68, 4.53)
    screenshot(
        s,
        "route_details.png",
        "Final route details",
        "Route cards · actual metrics\nTrade-off slider · avoided-street explanation",
        6.03,
        2.67,
        9.25,
        5.33,
    )
    notes(
        s,
        "Usability implemented in app/static/app.js and index.html: route comparison "
        "cards, metric glossary, avoided-street explanations, segment colours, "
        "comfort/time slider and chart, GPX export. Accessibility options include "
        "colour-vision palettes, increased text size, high contrast, reduced motion, "
        "dashed vs solid routes and browser speech synthesis. Selected factors only "
        "change route scoring; the metric cards retain all modelled factors. "
        "This is route planning with informed selection; no clinical validation or "
        "built-in live navigation is claimed. Screenshot is deliberately pending.",
    )

    s = new_slide(
        prs,
        9,
        "Implementation / integrated prototype",
        "What we actually built.",
        "A working code path from city data to a cyclist's route choice.",
    )
    text(s, "OUR IMPLEMENTATION", M, 2.35, 8.9, 0.37, 18, GREEN, True)
    modules = [
        ("Browser interface", "HTML / CSS / JS · comparison · slider · GPX export"),
        (
            "API + application engine",
            "Validation · orchestration · route metrics · explanations",
        ),
        (
            "Routing + environment model",
            "Profile costs · fuzzy discomfort · directed path search",
        ),
        (
            "GIS preprocessing + cache",
            "Graph export · height raster · shade tables · LRU cache",
        ),
    ]
    for i, (title, detail) in enumerate(modules):
        y = 2.96 + i * 1.0
        card(s, M, y, 8.91, 0.84)
        text(s, f"{i + 1:02d}", 0.97, y + 0.2, 0.62, 0.41, 23, GREEN, True)
        text(s, title, 1.85, y + 0.1, 7.38, 0.35, 21, INK, True)
        text(s, detail, 1.85, y + 0.49, 7.38, 0.3, 16.5, MUTED)
        if i < 3:
            line(s, 1.26, y + 0.84, 1.26, y + 0.98, GREEN, 1.2, arrow=True)
    card(s, 10.02, 2.96, 5.26, 1.86)
    text(s, "THIRD-PARTY LIBRARIES", 10.3, 3.17, 4.7, 0.34, 17, GREEN, True)
    text(
        s,
        "Leaflet · Flask · Pydantic\nSciPy · pvlib · scikit-fuzzy\nGeoPandas · Rasterio · OSMnx",
        10.3,
        3.7,
        4.7,
        1.04,
        20,
        INK,
    )
    card(s, 10.02, 5.03, 5.26, 1.86)
    text(s, "EXTERNAL APIs / DATA", 10.3, 5.24, 4.7, 0.34, 17, GREEN, True)
    text(
        s,
        "OpenStreetMap · GUGiK LoD1\nMSIP green cover · Open-Meteo\nGIOŚ monitoring stations",
        10.3,
        5.77,
        4.7,
        1.04,
        20,
        INK,
    )
    shape(s, M, 7.21, 14.56, 1.03, MINT, radius=True)
    text(s, "232", 1.01, 7.36, 1.2, 0.62, 36, GREEN, True)
    text(s, "tests passed locally", 2.42, 7.54, 3.78, 0.43, 21, INK, True)
    line(s, 6.31, 7.4, 6.31, 8.03, BORDER, 1.1)
    text(
        s,
        "Real-module integration verified on fixture graphs.\nFinal city capture needs the processed graph + shade artifacts.",
        6.64,
        7.46,
        8.31,
        0.7,
        17,
        MUTED,
    )
    notes(
        s,
        "Verified in this workspace: python -m pytest -q returned 232 passed, 5 skipped, "
        "122 subtests passed. The skipped checks require absent city graph/shade artifacts. "
        "The full real-module integration chain is tested with small fixture networks in "
        "tests/test_integration.py; those are not measured city routes. app/providers.py "
        "makes modules injectable and discloses mock/fallback status through /api/health. "
        "Offline geometry: OSM filtering including bicycle contraflow, deterministic eid "
        "indexing, CRS transformations, height raster and shade LUTs with graph fingerprints. "
        "Online: cached conditions, cost LRU, vector exposure and SciPy route search. "
        "Live refresh and stale fallback are implemented; recorded heatwave/smog days "
        "exist as 24-hour JSON scenarios with spatial weather grids. "
        "Optional pipeline/heat_map_build.py implements Landsat 8/9 surface-anomaly "
        "processing via Microsoft Planetary Computer; edge_heat.npz is absent, "
        "so its correction is not active here. No machine-learned microclimate model, "
        "NSGA-II, streaming evolutionary search, or reproducible city benchmark is claimed.",
    )

    s = new_slide(
        prs, 10, "One product / connected value", "One product. Two impacts.", dark=True
    )
    card(s, M, 2.4, 7.1, 3.6, dark=True)
    card(s, 8.17, 2.4, 7.11, 3.6, dark=True)
    text(s, "SPORT & HEALTHCARE", 1.06, 2.74, 6.41, 0.39, 18, LIME, True)
    text(
        s, "Make outdoor rides\nmore informed.", 1.06, 3.35, 6.41, 1.08, 30, WHITE, True
    )
    text(
        s,
        "Support cycling and active travel\nUnderstand air, heat, UV and shade\nChoose conditions that suit the ride",
        1.06,
        4.91,
        6.41,
        1.01,
        20,
        PALE,
    )
    text(s, "SMART CITY", 8.51, 2.74, 6.41, 0.39, 18, LIME, True)
    text(
        s,
        "Turn urban data\ninto a daily decision.",
        8.51,
        3.35,
        6.41,
        1.08,
        30,
        WHITE,
        True,
    )
    text(
        s,
        "Plan bicycle journeys on the road network\nConnect geometry with environmental data\nUse city infrastructure with context",
        8.51,
        4.91,
        6.41,
        1.01,
        20,
        PALE,
    )
    text(s, "BiKing gives every journey", M, 6.65, 14.56, 0.64, 36, WHITE, True)
    text(s, "the context of the city around it.", M, 7.34, 14.56, 0.64, 36, LIME, True)
    notes(
        s,
        "The final category mapping expresses intended practical value: support "
        "cycling/outdoor activity and informed route selection; turn urban geometry, "
        "the bicycle network and environmental data into individual journey decisions. "
        "No user study, increase in cycling, citywide benefit measurement, disease "
        "prevention, medical effect or guaranteed exposure reduction is claimed. "
        "Closing line connects both categories through the same product story.",
    )
    return prs


def write_checklist() -> None:
    lines = [
        "# BiKing — final screenshots",
        "",
        (
            "The finished presentation already exists at `BiKing_HackYeah.pptx`. "
            "Missing images have designed placeholders; all slide diagrams are editable."
        ),
        "",
        (
            "Put final screenshots in `presentation/assets/screenshots/`. "
            "The generator automatically contains each image in its existing frame without "
            "cropping. Existing screenshots are not modified."
        ),
        "",
        "## Capture readiness",
        "",
        (
            "This checkout contains the implemented Flask/Leaflet app, environmental "
            "scenarios, fuzzy tables and a synthetic shade preview. It currently lacks the "
            "city-wide `graph.npz`, `edges.parquet`, `edge_coords.npz`, `shade.npy`, "
            "`shade_bins.json` and `edge_tree_frac.npy` in `data/processed/`. "
            "Load matching artifacts before capturing real Kraków routing. "
            "The optional `edge_heat.npz` is also absent."
        ),
        "",
        (
            "Check `/api/health` before capture. A mock fallback must not be presented "
            "as real city data. Historical scenario data is legitimate when the selected "
            "scenario and departure time remain visible. Do not replace missing screenshots "
            "with generated UI or hand-drawn routes."
        ),
        "",
    ]
    for spec in SCREENSHOT_SPECS:
        status = (
            "Present"
            if (SCREENSHOTS / spec["filename"]).is_file()
            else "Missing — placeholder"
        )
        lines += [
            f"## Slide {spec['slide']} — {spec['filename']}",
            "",
            f"- Status: {status}.",
            f"- Visible: {spec['visible']}",
            f"- Recommended application state: {spec['state']}",
            f"- Must not be visible: {spec['exclude']}",
            f"- Why it matters: {spec['why']}",
            f"- Recommended size: {spec['size']}",
            "",
        ]
    lines += [
        "## Slide 7 metric cards",
        "",
        (
            "All current values are `—`. After capturing the final comparison, use "
            "the exact same `/api/route` response to enter both routes' `time_min`, "
            "`shade_pct`, `utci_avg_c` and `pm25_dose_ug` in the editable cards. "
            "BiKing in these cards means More comfortable. Preserve the model "
            "labels for UTCI and PM2.5 dose. Screenshot replacement does not invent "
            "or infer metric values."
        ),
        "",
        "## Shadow screenshot",
        "",
        (
            "No `shadow_debug.png` is required. Slide 6 is complete with a native "
            "PowerPoint geometry diagram and the actual ray-marching inequality. "
            "The existing `data/shade_preview/shade_comparison.png` is synthetic and "
            "is deliberately not used as a real-city screenshot."
        ),
        "",
    ]
    (OUT / "SCREENSHOTS_NEEDED.md").write_text("\n".join(lines), encoding="utf-8")


def preview_font(size_pt, bold=False):
    filename = "LiberationSans-Bold.ttf" if bold else "LiberationSans-Regular.ttf"
    candidates = [
        FONT_DIR / filename,
        Path("/usr/share/fonts/truetype/liberation2") / filename,
    ]
    for candidate in candidates:
        if candidate.is_file():
            return ImageFont.truetype(
                str(candidate), max(1, round(size_pt * SCALE / 72))
            )
    return ImageFont.truetype("DejaVuSans.ttf", max(1, round(size_pt * SCALE / 72)))


def solid_color(fill, default=None):
    try:
        return "#" + str(fill.fore_color.rgb) if fill.type else default
    except (AttributeError, TypeError):
        return default


def preview_and_validate(prs: Presentation) -> dict:
    """Inspect the saved shapes using Pillow; this is not an office rendering."""
    errors, slides, text_count = [], [], 0
    for num, slide in enumerate(prs.slides, 1):
        text_rects = []
        bg = solid_color(slide.background.fill, "#" + BG)
        canvas = Image.new("RGB", (round(W * SCALE), round(H * SCALE)), bg)
        draw = ImageDraw.Draw(canvas)
        for obj in slide.shapes:
            x, y, w, h = [
                float(v) / EMU * SCALE
                for v in (obj.left, obj.top, obj.width, obj.height)
            ]
            if min(x, y) < -1 or x + w > W * SCALE + 1 or y + h > H * SCALE + 1:
                errors.append(f"Slide {num}: object outside slide: {obj.name}")
            if obj.shape_type == MSO_SHAPE_TYPE.PICTURE:
                im = Image.open(io.BytesIO(obj.image.blob)).convert("RGBA")
                im = im.resize(
                    (max(1, round(w)), max(1, round(h))), Image.Resampling.LANCZOS
                )
                canvas.paste(im, (round(x), round(y)), im)
                continue
            fill = solid_color(obj.fill) if hasattr(obj, "fill") else None
            stroke = None
            try:
                stroke = "#" + str(obj.line.color.rgb) if obj.line.fill.type else None
            except (AttributeError, TypeError):
                pass
            width = max(1, round((obj.line.width or Pt(1)) / EMU * SCALE))
            bounds = (round(x), round(y), round(x + w), round(y + h))
            if obj.shape_type == MSO_SHAPE_TYPE.LINE:
                transform = obj._element.find(".//{*}xfrm")
                flip_h = transform is not None and transform.get("flipH") == "1"
                flip_v = transform is not None and transform.get("flipV") == "1"
                a = (x + w if flip_h else x, y + h if flip_v else y)
                b = (x if flip_h else x + w, y if flip_v else y + h)
                dash = obj._element.find(".//{*}prstDash")
                if dash is not None:
                    distance = math.dist(a, b)
                    for start in range(0, max(1, round(distance)), 13):
                        end = min(start + 7, distance)
                        aa = tuple(
                            a[i] + (b[i] - a[i]) * start / max(distance, 1)
                            for i in range(2)
                        )
                        bb = tuple(
                            a[i] + (b[i] - a[i]) * end / max(distance, 1)
                            for i in range(2)
                        )
                        draw.line([aa, bb], fill=stroke or "#" + BORDER, width=width)
                else:
                    draw.line([a, b], fill=stroke or "#" + BORDER, width=width)
                arrow = obj._element.find(".//{*}tailEnd")
                if arrow is not None:
                    angle = math.atan2(b[1] - a[1], b[0] - a[0])
                    points = [
                        b,
                        (
                            b[0] - 9 * math.cos(angle - 0.43),
                            b[1] - 9 * math.sin(angle - 0.43),
                        ),
                        (
                            b[0] - 9 * math.cos(angle + 0.43),
                            b[1] - 9 * math.sin(angle + 0.43),
                        ),
                    ]
                    draw.polygon(points, fill=stroke or "#" + BORDER)
                continue
            if obj.shape_type == MSO_SHAPE_TYPE.FREEFORM:
                path = obj._element.find(".//{*}path")
                if path is not None:
                    points = [
                        (
                            x + int(pt.get("x")) / int(path.get("w")) * w,
                            y + int(pt.get("y")) / int(path.get("h")) * h,
                        )
                        for pt in path.findall(".//{*}pt")
                    ]
                    draw.polygon(points, fill=fill)
            elif obj.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                preset = obj.auto_shape_type
                if preset == MSO_SHAPE.OVAL:
                    draw.ellipse(bounds, fill=fill, outline=stroke, width=width)
                elif preset == MSO_SHAPE.ROUNDED_RECTANGLE:
                    radius = min(w, h) * 0.12
                    draw.rounded_rectangle(
                        bounds, radius=radius, fill=fill, outline=stroke, width=width
                    )
                else:
                    draw.rectangle(bounds, fill=fill, outline=stroke, width=width)
            if not obj.has_text_frame or not obj.text:
                continue
            text_count += 1
            frame = obj.text_frame
            tx = x + frame.margin_left / EMU * SCALE
            ty = y + frame.margin_top / EMU * SCALE
            usable_w = w - (frame.margin_left + frame.margin_right) / EMU * SCALE
            usable_h = h - (frame.margin_top + frame.margin_bottom) / EMU * SCALE
            consumed = 0.0
            for p in frame.paragraphs:
                if not p.runs:
                    continue
                run = p.runs[0]
                size = run.font.size.pt
                font = preview_font(size, run.font.bold)
                value = p.text
                length = draw.textlength(value, font=font)
                if length > usable_w + 1:
                    errors.append(
                        f"Slide {num}: text too wide: {value!r} ({length:.1f}>{usable_w:.1f})"
                    )
                step = p.line_spacing / EMU * SCALE
                if p.alignment == PP_ALIGN.CENTER:
                    px = tx + (usable_w - length) / 2
                elif p.alignment == PP_ALIGN.RIGHT:
                    px = tx + usable_w - length
                else:
                    px = tx
                draw.text(
                    (px, ty + consumed),
                    value,
                    font=font,
                    fill="#" + str(run.font.color.rgb),
                    anchor="lt",
                )
                rect = draw.textbbox((px, ty + consumed), value, font=font, anchor="lt")
                text_rects.append((rect, value))
                consumed += step
            if consumed > usable_h + 2:
                errors.append(
                    f"Slide {num}: text box too short: {obj.text!r} ({consumed:.1f}>{usable_h:.1f})"
                )
        for i, (a, label_a) in enumerate(text_rects):
            for b, label_b in text_rects[i + 1 :]:
                if (
                    min(a[2], b[2]) - max(a[0], b[0]) > 1
                    and min(a[3], b[3]) - max(a[1], b[1]) > 1
                ):
                    errors.append(
                        f"Slide {num}: overlapping text: {label_a!r} / {label_b!r}"
                    )
        path = PREVIEWS / f"slide_{num:02d}.png"
        canvas.save(path)
        slides.append(canvas)
    thumbnail_w, thumbnail_h = 640, 360
    contact = Image.new(
        "RGB", (thumbnail_w * 2 + 48, 5 * (thumbnail_h + 40) + 70), "#DDE4DC"
    )
    d = ImageDraw.Draw(contact)
    label_font = preview_font(17, True)
    d.text(
        (20, 14),
        "BiKing / STATIC LAYOUT PREVIEWS / not office-rendered",
        font=label_font,
        fill="#" + INK,
    )
    for i, canvas in enumerate(slides):
        x = 16 + (i % 2) * (thumbnail_w + 16)
        y = 65 + (i // 2) * (thumbnail_h + 40)
        contact.paste(
            canvas.resize((thumbnail_w, thumbnail_h), Image.Resampling.LANCZOS), (x, y)
        )
        d.text(
            (x + 4, y + thumbnail_h + 7),
            f"{i + 1:02d}",
            font=preview_font(13, True),
            fill="#" + MUTED,
        )
    contact.save(PREVIEWS / "contact_sheet.png")
    report = {
        "file": str(PPTX.relative_to(ROOT)),
        "slides": len(prs.slides),
        "aspect_ratio": "16:9",
        "text_boxes_checked": text_count,
        "layout_errors": errors,
        "missing_screenshots": [
            spec["filename"]
            for spec in SCREENSHOT_SPECS
            if not (SCREENSHOTS / spec["filename"]).is_file()
        ],
        "office_rendered": False,
        "inspection_method": "Saved PPTX reopened; shape bounds and Arial-compatible text "
        "measurements checked; Pillow static layout previews generated.",
    }
    (OUT / "validation.json").write_text(
        json.dumps(report, indent=2) + "\n", encoding="utf-8"
    )
    if errors:
        raise ValueError("\n".join(errors))
    return report


def main() -> None:
    for directory in (OUT, SCREENSHOTS, PREVIEWS):
        directory.mkdir(parents=True, exist_ok=True)
    (SCREENSHOTS / ".gitkeep").touch()
    prs = make_deck()
    prs.save(PPTX)
    with ZipFile(PPTX) as package:
        if package.testzip() is not None:
            raise ValueError("Invalid PPTX archive")
    report = preview_and_validate(Presentation(PPTX))
    write_checklist()
    print(json.dumps(report, indent=2))
    print(f"Finished: {PPTX} ({PPTX.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
